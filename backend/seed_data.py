from pathlib import Path
import openpyxl
from backend.database import get_db_connection, init_db

BASE_DIR = Path(__file__).resolve().parent.parent


def get_seed_file() -> Path:
    candidates = [
        BASE_DIR / "data" / "events_seed.xlsx",
        BASE_DIR / "zenith26-fest-system" / "data" / "events_seed.xlsx",
        BASE_DIR / "events_seed.xlsx",
        Path.cwd() / "data" / "events_seed.xlsx",
        Path.cwd() / "events_seed.xlsx",
    ]
    for c in candidates:
        if c.exists():
            return c
    return BASE_DIR / "data" / "events_seed.xlsx"


def seed_events():
    """Seed the events table from data/events_seed.xlsx if not already seeded."""
    init_db()

    conn = get_db_connection()
    cursor = conn.cursor()

    # Skip seeding if the events table already has rows
    cursor.execute("SELECT COUNT(*) FROM events")
    count = cursor.fetchone()[0]
    if count > 0:
        print(f"Database already contains {count} events. Skipping seeding.")
        conn.close()
        return

    seed_file = get_seed_file()
    # Check if the spreadsheet exists
    if not seed_file.exists():
        conn.close()
        raise FileNotFoundError(
            f"Seed data file not found at '{seed_file}'. Please copy the provided events_seed.xlsx file into the data/ folder before running this script."
        )

    wb = openpyxl.load_workbook(str(seed_file))
    if "Events" not in wb.sheetnames:
        conn.close()
        raise ValueError(
            f"Sheet 'Events' not found in '{SEED_FILE}'. Expected an 'Events' sheet."
        )

    sheet = wb["Events"]

    # Header check
    expected_headers = [
        "id", "title", "category", "description", "event_date",
        "event_time", "venue", "capacity", "fee", "icon"
    ]

    header_row = [cell.value for cell in sheet[1]]
    if [h.lower() if isinstance(h, str) else h for h in header_row[:10]] != expected_headers:
        print(f"Warning: Header row does not strictly match expected headers: {header_row}")

    rows_inserted = 0
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue

        (
            event_id,
            title,
            category,
            description,
            event_date,
            event_time,
            venue,
            capacity,
            fee,
            icon,
        ) = row[:10]

        cursor.execute(
            """
            INSERT INTO events (id, title, category, description, event_date, event_time, venue, capacity, fee, icon)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                int(event_id),
                str(title),
                str(category),
                str(description) if description is not None else "",
                str(event_date),
                str(event_time),
                str(venue),
                int(capacity),
                int(fee) if fee is not None else 0,
                str(icon) if icon is not None else "",
            ),
        )
        rows_inserted += 1

    conn.commit()
    conn.close()
    print(f"Successfully seeded {rows_inserted} events from {SEED_FILE.name}.")


if __name__ == "__main__":
    seed_events()

    # Display seeded events to confirm
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT id, title, category, event_date, event_time, venue, capacity, fee, icon FROM events ORDER BY id ASC")
    rows = cursor.fetchall()
    print(f"\n--- Events in fest.db ({len(rows)} total) ---")
    for r in rows:
        print(f"#{r['id']}: [{r['icon']} {r['category']}] {r['title']} | {r['event_date']} {r['event_time']} @ {r['venue']} | Cap: {r['capacity']} | Fee: ₹{r['fee']}")
    conn.close()
